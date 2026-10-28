import json, random, re, sys, time
from pathlib import Path

from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

TOTAL_USUARIOS = 200
NUM = r"([\d.,]+(?:\s*(?:mil|M|K))?)"
RE_SEGUIDORES = re.compile(NUM + r"\s*(?:seguidores|followers)\b", re.I)
RE_SEGUIDOS = re.compile(NUM + r"\s*(?:seguidos|following)\b", re.I)
# elemento más pequeño que contenga el nombre del usuario y las etiquetas de la tarjeta
JS_TARJETA = """(u) => { u = u.toLowerCase(); let best = null;
  for (const el of document.querySelectorAll('div')) {
    const t = el.innerText; if (!t || t.length > 500) continue; const l = t.toLowerCase();
    if (l.includes(u) && /(publicaciones|posts)/.test(l) && /(seguidores|followers)/.test(l)
        && /(seguidos|following)/.test(l) && (!best || t.length < best.length)) best = t; }
  return best; }"""


def leer_username_de_scraper():
    try:
        m = re.search(r'^USERNAME\s*=\s*["\']([\w.]+)["\']', Path("scraper.py").read_text(encoding="utf-8"), re.M)
    except FileNotFoundError:
        m = None
    return m.group(1) if m else "sausage_player"


def convertir_cookies(texto):
    ss = {"no_restriction": "None", "lax": "Lax", "strict": "Strict"}
    return [{"name": c["name"], "value": c["value"], "domain": c["domain"], "path": c.get("path", "/"),
             "secure": c.get("secure", False), "httpOnly": c.get("httpOnly", False),
             "sameSite": ss.get(str(c.get("sameSite", "")).lower(), "Lax"),
             **({"expires": int(c["expirationDate"])} if not c.get("session") and "expirationDate" in c else {})}
            for c in json.loads(texto)]


def pausa(a=1.0, b=2.5):
    time.sleep(random.uniform(a, b))


def leer_tarjeta(page, usuario, espera=4.0):
    """Espera la tarjeta del usuario y devuelve (seguidores, seguidos) o (None, None)."""
    fin = time.time() + espera
    while time.time() < fin:
        t = page.evaluate(JS_TARJETA, usuario) or ""
        ms, mf = RE_SEGUIDORES.search(t), RE_SEGUIDOS.search(t)
        if ms and mf:
            return ms.group(1).strip(), mf.group(1).strip()
        time.sleep(0.4)
    return None, None

def abrir_lista_seguidores(page, cuenta):
    page.goto(f"https://www.instagram.com/{cuenta}/", wait_until="domcontentloaded")
    pausa(3, 5)
    if "accounts/login" in page.url:
        raise SystemExit("Cookies inválidas o expiradas. Vuelve a exportarlas.")
    page.wait_for_selector("header", timeout=15000)
    enlace = page.locator('a[href*="/followers"]')
    if enlace.count() == 0:
        enlace = page.get_by_role("link", name=re.compile(r"seguidores|followers", re.I))
    if enlace.count() == 0:
        raise SystemExit("⚠️ No se encontró el enlace de seguidores (¿cuenta privada, cookies inválidas "
                         "o Instagram cambió la página?).")
    enlace.first.click()
    page.wait_for_selector('div[role="dialog"]', timeout=15000)
    pausa(2, 3)


def usuarios_visibles(page, cuenta):
    """Usuarios visibles en la lista (el enlace cuyo texto es igual al usuario es el nombre)."""
    anclas = page.eval_on_selector_all('div[role="dialog"] a[href^="/"]',
                                       "els => els.map(e => [e.getAttribute('href'), e.innerText.trim()])")
    return list(dict.fromkeys(h.strip("/") for h, t in anclas
                              if re.fullmatch(r"/[\w.]+/", h or "") and t == h.strip("/") and t != cuenta))


def desplazar_lista(page):
    caja = page.locator('div[role="dialog"]').first.bounding_box()
    if caja:
        page.mouse.move(caja["x"] + caja["width"] / 2, caja["y"] + caja["height"] / 2)
    page.mouse.wheel(0, random.randint(500, 800))
    pausa(1.2, 2.2)


def _num(v):
    """'2593' -> 2593 | '1.234' -> 1234 | '1,2 mil' queda como texto."""
    t = v.replace(".", "")
    return int(t) if t.isdigit() else v


def guardar(datos, base):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    Path(f"{base}.json").write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = "Usuarios"
    ws.append(["Usuario", "Seguidores", "Seguidos"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F2340")
    for d in datos:
        ws.append([d["usuario"], _num(d["seguidores"]), _num(d["seguidos"])])
    for col, ancho in zip("ABC", (32, 14, 14)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    wb.save(f"{base}.xlsx")


def cargar_previo(base):
    """Permite reanudar si la ejecución se interrumpió."""
    try:
        return json.loads(Path(f"{base}.json").read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return []


def recolectar_usuarios(page, cuenta):
    """Recolecta TOTAL_USUARIOS con una página de Playwright ya abierta y con sesión
    (scraper.py la llama con su propio navegador, así no se abre un segundo)."""
    base = f"{cuenta}_usuarios_{TOTAL_USUARIOS}"
    datos = cargar_previo(base)
    hechos = {d["usuario"] for d in datos}
    if datos:
        print(f"Reanudando: ya hay {len(datos)} usuarios en '{base}.json'.")
    if len(datos) < TOTAL_USUARIOS:
        print(f"Abriendo la lista de seguidores de @{cuenta}...")
        abrir_lista_seguidores(page, cuenta)

    vacias = 0
    while len(datos) < TOTAL_USUARIOS and vacias < 8:
        nuevos = [u for u in usuarios_visibles(page, cuenta) if u not in hechos]
        vacias = 0 if nuevos else vacias + 1
        for u in nuevos[:TOTAL_USUARIOS - len(datos)]:
            hechos.add(u)
            enlace = page.locator(f'div[role="dialog"] a[href="/{u}/"]').filter(has_text=u).first
            try:
                enlace.scroll_into_view_if_needed(timeout=5000)
                enlace.hover(timeout=5000)
            except PWTimeout:
                print(f"  · @{u}: no se pudo pasar el mouse, se omite.")
                continue
            seguidores, seguidos = leer_tarjeta(page, u)
            if seguidores:
                datos.append({"usuario": u, "seguidores": seguidores, "seguidos": seguidos})
                guardar(datos, base)   # guarda en cada paso por si se interrumpe
                print(f"[{len(datos)}/{TOTAL_USUARIOS}] @{u}  seguidores={seguidores}  seguidos={seguidos}")
            else:
                print(f"  · @{u}: la tarjeta no apareció, se omite.")
            page.mouse.move(5, 5)      # saca el mouse para que la tarjeta se cierre
            pausa(0.8, 1.8)
        desplazar_lista(page)

    guardar(datos, base)
    if len(datos) < TOTAL_USUARIOS:
        print(f"⚠️ Solo se pudieron recolectar {len(datos)} de {TOTAL_USUARIOS} usuarios.")
    print(f"✅ Datos guardados en '{base}.json' y '{base}.xlsx'.")
    return datos


def main():
    """Uso independiente: abre su propio navegador."""
    cuenta = sys.argv[1] if len(sys.argv) > 1 else leer_username_de_scraper()
    try:
        cookies = Path("cookies.json").read_text(encoding="utf-8")
    except FileNotFoundError:
        raise SystemExit("⚠️ Error: No se encontró el archivo 'cookies.json' en la carpeta del proyecto.")

    base = f"{cuenta}_usuarios_{TOTAL_USUARIOS}"
    previo = cargar_previo(base)
    if len(previo) >= TOTAL_USUARIOS:   # ya está todo: solo regenera los archivos, sin navegador
        guardar(previo, base)
        return print(f"✅ Ya había {len(previo)} usuarios. Archivos regenerados.")

    with sync_playwright() as p:
        navegador = p.chromium.launch(headless=False)   # visible: ves la macro trabajando
        contexto = navegador.new_context(
            viewport={"width": 1280, "height": 900}, locale="es-ES",
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                       "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36")
        contexto.add_cookies(convertir_cookies(cookies))
        recolectar_usuarios(contexto.new_page(), cuenta)
        navegador.close()


if __name__ == "__main__":
    main()