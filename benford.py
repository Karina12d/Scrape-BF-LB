import json, math, re, sys
from collections import Counter
from pathlib import Path

import matplotlib
matplotlib.use("Agg")   # genera el PNG sin necesitar ventana
import matplotlib.pyplot as plt

from recolector_seguidores import TOTAL_USUARIOS, leer_username_de_scraper

CAMPO, GL, MIN_MUESTRA = "seguidos", 8, 50
CRITICO_5, CRITICO_1 = 15.507, 20.090   # chi-cuadrado con 8 gl: p = 0.05 y p = 0.01
COLORES = {"Humano": "#10B981", "Posible bot": "#F59E0B", "Bot": "#EF4444"}
D = range(1, 10)


def benford(d):
    return math.log10(1 + 1 / d)


def parsear_cantidad(texto):
    """'2593' -> 2593 | '2.593' -> 2593 | '1,2 mil' -> 1200 | '3,4 M' -> 3400000 | inválido -> None"""
    m = re.fullmatch(r"([\d.,]+)\s*(mil|k|m)?", str(texto).strip().lower())
    if not m:
        return None
    numero, sufijo = m.groups()
    if sufijo:   # con sufijo el separador es decimal; sin sufijo, de miles
        return round(float(numero.replace(",", ".")) * {"mil": 1e3, "k": 1e3, "m": 1e6}[sufijo])
    return int(re.sub(r"[.,]", "", numero))


def analizar(valores):
    """Conteo del primer dígito (los ceros se ignoran) y chi-cuadrado contra Benford."""
    c = Counter(int(str(v)[0]) for v in valores if v)
    conteo = {d: c[d] for d in D}
    n = sum(conteo.values())
    esp = {d: n * benford(d) for d in D}
    contrib = {d: (conteo[d] - esp[d]) ** 2 / esp[d] for d in D}
    chi = sum(contrib.values())
    x = chi / 2   # p-valor exacto para gl par: e^(-x) * sum(x^k / k!), k < gl/2
    p = math.exp(-x) * sum(x ** k / math.factorial(k) for k in range(GL // 2))
    v = "Humano" if chi <= CRITICO_5 else "Posible bot" if chi <= CRITICO_1 else "Bot"
    return {"conteo": conteo, "n": n, "esp": esp, "contrib": contrib, "chi": chi, "p": p, "veredicto": v}


def graficar(r, cuenta, archivo_png):
    """Una sola gráfica: barras observadas + curva de Benford + curva observada + área de distancia."""
    n, chi, v, conteo = r["n"], r["chi"], r["veredicto"], r["conteo"]
    xs = list(D)
    obs = [100 * conteo[d] / n for d in xs]
    esp = [100 * benford(d) for d in xs]
    fig = plt.figure(figsize=(11.5, 7.6))
    ax, escala = fig.add_gridspec(2, 1, height_ratios=[5, 1.15], hspace=0.42).subplots()

    ax.bar(xs, obs, width=0.7, color="#FACC15", alpha=0.9, label="Observado (barras)", zorder=1)
    ax.fill_between(xs, obs, esp, interpolate=True, facecolor="#3B4FD8", alpha=0.22, hatch="///",
                    edgecolor="#3B4FD8", linewidth=0, label="Distancia (área entre curvas)")
    ax.plot(xs, esp, "o-", color="#EF4444", linewidth=2.8, markersize=7, label="Ley de Benford (esperado)", zorder=3)
    ax.plot(xs, obs, "o-", color="#84B30F", linewidth=2.8, markersize=7, label="Observado – chi-cuadrado", zorder=4)
    for d, o, e in zip(xs, obs, esp):
        ax.text(d, max(o, e) + 1, conteo[d], ha="center", fontsize=10, fontweight="bold", color="#1F2340")
    ax.set(xticks=xs, ylim=(0, max(obs + esp) + 14), ylabel="Frecuencia (%)",
           xticklabels=[f"{d}\nχ²ᵢ = {r['contrib'][d]:.2f}" for d in xs],
           xlabel="Primer dígito   (χ²ᵢ = (O − E)² / E: aporte de cada dígito a la distancia)",
           title=f"Ley de Benford y chi-cuadrado – {CAMPO} de los usuarios de @{cuenta} (n = {n})")
    ax.title.set(fontsize=13, fontweight="bold")
    ax.grid(alpha=0.3)
    ax.legend(loc="upper center", ncol=2, fontsize=9.5)
    ax.text(0.985, 0.80, f"Distancia χ² = Σ(O−E)²/E = {chi:.2f}\ngl = {GL}   p-valor = {r['p']:.4f}\n"
            f"RESULTADO: {v.upper()}", transform=ax.transAxes, ha="right", va="top", fontsize=11,
            fontweight="bold", color=COLORES[v], bbox=dict(boxstyle="round,pad=0.5", fc="white", ec=COLORES[v], lw=2))

    # escala de decisión: dónde cae la distancia χ²
    for ini, fin, z in ((0, CRITICO_5, "Humano"), (CRITICO_5, CRITICO_1, "Posible bot"), (CRITICO_1, 35, "Bot")):
        escala.axvspan(ini, fin, color=COLORES[z], alpha=0.35)
        escala.text((ini + fin) / 2, 0.5, z.upper(), ha="center", va="center", fontsize=9)
    x = min(chi, 34)
    escala.axvline(x, color=COLORES[v], linewidth=4)
    escala.annotate(f"χ² = {chi:.2f}", (x, 1), (0, 6), textcoords="offset points", ha="center",
                    fontsize=10, fontweight="bold", color=COLORES[v], annotation_clip=False)
    escala.set(xlim=(0, 35), ylim=(0, 1), yticks=[], xlabel=f"Escala de decisión según la distancia χ² ({GL} grados de libertad)")
    fig.savefig(archivo_png, dpi=150, bbox_inches="tight")
    plt.close(fig)


def guardar_en_excel(archivo_json, r):
    """Escribe el conteo y el chi-cuadrado en la MISMA hoja del Excel, desde la columna E."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    ruta = re.sub(r"\.json$", ".xlsx", archivo_json)
    try:
        wb = load_workbook(ruta)
    except FileNotFoundError:
        return print(f"⚠️ No existe '{ruta}'. Ejecuta primero recolector_seguidores.py.")
    ws = wb["Usuarios"] if "Usuarios" in wb.sheetnames else wb.active
    n = r["n"]
    filas = [["Primer dígito", f"Conteo {CAMPO}", "Observado %", "Benford %", "Conteo esperado", "(O-E)²/E"]]
    filas += [[d, r["conteo"][d], r["conteo"][d] / n, benford(d), r["esp"][d], r["contrib"][d]] for d in D]
    filas += [["Total", n], [], ["Chi-cuadrado (χ²)", round(r["chi"], 3)], ["Grados de libertad", GL],
              ["Valor crítico 5%", CRITICO_5], ["Valor crítico 1%", CRITICO_1],
              ["p-valor", round(r["p"], 4)], ["Resultado", r["veredicto"].upper()]]
    for i, fila in enumerate(filas, 1):
        for j, valor in enumerate(fila, 5):
            ws.cell(i, j, valor)
    for i in range(2, 11):
        ws.cell(i, 7).number_format = ws.cell(i, 8).number_format = "0.0%"
    for c in ws[1][4:10]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="7C3AED")
    for col in "EFGHIJ":
        ws.column_dimensions[col].width = 18
    res = ws.cell(len(filas), 6)
    res.font, res.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=COLORES[r["veredicto"]][1:])
    try:
        wb.save(ruta)
        print(f"✅ Tabla de Benford y chi-cuadrado agregada en '{ruta}' (hoja '{ws.title}', desde la columna E).")
    except PermissionError:
        print(f"⚠️ No se pudo guardar '{ruta}': ciérralo en Excel y vuelve a ejecutar benford.py.")


def main():
    if len(sys.argv) > 1:
        archivo = sys.argv[1]
        cuenta = re.sub(r"_usuarios_\d+\.json$", "", re.split(r"[\\/]", archivo)[-1])
    else:
        cuenta = leer_username_de_scraper()
        archivo = f"{cuenta}_usuarios_{TOTAL_USUARIOS}.json"
    try:
        usuarios = json.loads(Path(archivo).read_text(encoding="utf-8"))
    except FileNotFoundError:
        raise SystemExit(f"⚠️ No se encontró '{archivo}'. Ejecuta primero recolector_seguidores.py.")
    valores = [parsear_cantidad(u.get(CAMPO)) for u in usuarios]
    if not any(valores):
        raise SystemExit("⚠️ No hay valores válidos para analizar.")

    r = analizar(valores)
    print(f"Usuarios leídos: {len(usuarios)} | valores con primer dígito válido: {r['n']}\n")
    print("dígito | conteo | observado % | Benford %")
    for d in D:
        print(f"   {d}   | {r['conteo'][d]:>6} | {100 * r['conteo'][d] / r['n']:>10.1f}  | {100 * benford(d):>8.1f}")
    print(f"\nχ² = {r['chi']:.3f} (gl = {GL}, crítico 5% = {CRITICO_5}, 1% = {CRITICO_1})  p = {r['p']:.4f}")
    print(f"RESULTADO: {r['veredicto'].upper()}")
    if r["n"] < MIN_MUESTRA:
        print(f"⚠️ Solo {r['n']} valores válidos (< {MIN_MUESTRA}): el resultado no es confiable.")

    png = f"benford_{cuenta}.png"
    graficar(r, cuenta, png)
    print(f"\n✅ Gráfica guardada en '{png}'.")
    guardar_en_excel(archivo, r)


if __name__ == "__main__":
    main()